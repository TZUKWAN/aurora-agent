#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
insert_ppt_images.py
将PPT导出的PNG图片插入到49份DOCX商业计划书中
"""
import json
import os
import re
import time
from pathlib import Path
from PIL import Image as PILImage
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = Path(r"D:\计划书AI")
OUTPUT_DIR = BASE_DIR / "output_v2"
IMAGE_BASE = OUTPUT_DIR / "images"
PROJECTS_PATH = OUTPUT_DIR / "projects_list.json"

# 图号文件名映射: 图1-1 -> 图1_1.png
FIG_FILENAME_MAP = {
    "图1-1": "图1_1.png", "图1-2": "图1_2.png", "图1-3": "图1_3.png",
    "图1-4": "图1_4.png", "图1-5": "图1_5.png", "图1-6": "图1_6.png",
    "图1-7": "图1_7.png", "图1-8": "图1_8.png", "图1-9": "图1_9.png",
    "图2-1": "图2_1.png", "图2-2": "图2_2.png", "图2-3": "图2_3.png",
    "图2-4": "图2_4.png", "图2-5": "图2_5.png", "图2-6": "图2_6.png",
    "图2-7": "图2_7.png", "图2-8": "图2_8.png", "图2-9": "图2_9.png",
    "图2-10": "图2_10.png", "图2-14": "图2_14.png",
    "图2-11": "图2_11.png", "图2-12": "图2_12.png", "图2-13": "图2_13.png",
    "图2-15": "图2_15.png", "图2-16": "图2_16.png", "图2-17": "图2_17.png",
    "图2-18": "图2_18.png", "图2-19": "图2_19.png", "图2-20": "图2_20.png",
    "图2-21": "图2_21.png", "图2-52": "图2_52.png", "图2-53": "图2_53.png",
    "图2-54": "图2_54.png", "图2-55": "图2_55.png", "图2-56": "图2_56.png",
    "图2-57": "图2_57.png", "图2-58": "图2_58.png", "图2-59": "图2_59.png",
    "图2-60": "图2_60.png", "图2-61": "图2_61.png",
    "图3-1": "图3_1.png", "图3-2": "图3_2.png",
    "图4-1": "图4_1.png",
}

def find_docx(name):
    for f in os.listdir(str(OUTPUT_DIR)):
        if f.endswith('.docx') and not f.startswith('~') and name in f:
            return OUTPUT_DIR / f
    return None

def insert_images_for_project(project):
    idx = project['idx']
    name = project['name']
    img_dir = IMAGE_BASE / f"p{idx}"
    docx_path = find_docx(name)

    if not docx_path:
        print(f"  [{idx+1}] {name}: DOCX不存在")
        return 0, 0

    doc = Document(str(docx_path))
    paras = doc.paragraphs
    inserted = 0
    missing = 0

    for i, p in enumerate(paras):
        m = re.match(r'^(图\d+-\d+)\s', p.text.strip())
        if not m:
            continue
        fig_key = m.group(1)
        img_filename = FIG_FILENAME_MAP.get(fig_key)
        if not img_filename:
            continue
        img_path = img_dir / img_filename
        if not img_path.exists():
            missing += 1
            continue

        # 图注前的空段落作为图片位置
        if i > 0 and not paras[i - 1].text.strip():
            target = paras[i - 1]
            # 清除残留
            for child in list(target._element):
                tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                if tag == 'r':
                    target._element.remove(child)

            # 图片尺寸：竖屏小，横屏大
            try:
                w, h = PILImage.open(img_path).size
                width = Inches(2.8) if h > w * 1.5 else Inches(5.2)
            except Exception:
                width = Inches(5.2)

            run = target.add_run()
            try:
                run.add_picture(str(img_path), width=width)
                inserted += 1
            except Exception as e:
                print(f"    insert {fig_key} fail: {e}")

    # 居中图片段落
    for p in doc.paragraphs:
        if not p.text.strip():
            has_img = any(r._element.findall(qn('w:drawing')) for r in p.runs)
            if has_img:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(docx_path))
    return inserted, missing

def main():
    print("=" * 60)
    print(" PPT图表PNG插入DOCX")
    print("=" * 60)

    projects = json.loads(PROJECTS_PATH.read_text(encoding="utf-8"))
    total = len(projects)
    total_inserted = 0
    total_missing = 0

    t_start = time.time()
    for idx, project in enumerate(projects):
        t0 = time.time()
        inserted, missing = insert_images_for_project(project)
        elapsed = time.time() - t0
        total_inserted += inserted
        total_missing += missing
        print(f"[{idx+1}/{total}] {project['name']}: {inserted}张插入, {missing}张缺失 ({elapsed:.1f}s)")

    total_elapsed = time.time() - t_start
    print(f"\n{'='*60}")
    print(f"完成: {total}个项目, 共插入 {total_inserted} 张图片, 缺失 {total_missing} 张")
    print(f"耗时 {total_elapsed:.1f}s")
    print(f"{'='*60}")

if __name__ == "__main__":
    main()
