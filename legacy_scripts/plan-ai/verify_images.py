from docx import Document
from pathlib import Path
import re

OUTPUT_DIR = Path(r'D:\计划书AI\output_v2')
docx_files = [f for f in OUTPUT_DIR.glob('*.docx') if not f.name.startswith('~$')]

for docx_path in docx_files[:5]:
    doc = Document(str(docx_path))
    
    # Count total images in document
    total_images = 0
    image_paras = []
    for i, p in enumerate(doc.paragraphs):
        has_img = any(r._element.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing') for r in p.runs)
        if has_img:
            total_images += 1
            image_paras.append(i)
    
    # Count figure captions
    fig_captions = sum(1 for p in doc.paragraphs if re.match(r'^图\d+-\d+\s', p.text.strip()))
    
    print(f'{docx_path.name[:25]}...')
    print(f'  Images: {total_images}, Captions: {fig_captions}')
    print(f'  Image paragraphs: {image_paras}')
    
    # Check for images in caption paragraphs (shouldn't happen)
    bad = 0
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if re.match(r'^图\d+-\d+\s', text):
            has_img = any(r._element.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing') for r in p.runs)
            if has_img:
                bad += 1
    if bad:
        print(f'  WARNING: {bad} images inside caption paragraphs')
    else:
        print(f'  OK: No images inside caption paragraphs')
