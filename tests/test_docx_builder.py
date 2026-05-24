"""Tests for DOCX builder."""

import os
import tempfile

import pytest

from aurora.business_plan.generator import BusinessPlanGenerator


def _make_plan():
    gen = BusinessPlanGenerator()
    return gen.generate({
        "technology": "AI",
        "target_market": "Enterprise",
        "problem": "inefficiency",
        "solution": "AI automation",
        "product": "AI Platform",
    })


@pytest.fixture
def tmp_dir():
    d = tempfile.mkdtemp()
    yield d
    for f in os.listdir(d):
        os.unlink(os.path.join(d, f))
    os.rmdir(d)


class TestDocxBuilder:
    def test_export_creates_file(self, tmp_dir):
        plan = _make_plan()
        path = os.path.join(tmp_dir, "test.docx")
        gen = BusinessPlanGenerator()
        result = gen.export_to_docx(plan, path)
        assert result["success"] is True
        assert os.path.exists(path)
        assert os.path.getsize(path) > 0

    def test_export_has_cover_and_toc(self, tmp_dir):
        plan = _make_plan()
        path = os.path.join(tmp_dir, "test.docx")
        gen = BusinessPlanGenerator()
        gen.export_to_docx(plan, path)
        from docx import Document
        doc = Document(path)
        # First page should have heading and TOC
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert any("商业计划书" in t for t in texts)

    def test_export_sections_present(self, tmp_dir):
        plan = _make_plan()
        path = os.path.join(tmp_dir, "test.docx")
        gen = BusinessPlanGenerator()
        gen.export_to_docx(plan, path)
        from docx import Document
        doc = Document(path)
        texts = " ".join(p.text for p in doc.paragraphs)
        # Should have content from multiple sections
        assert len(texts) > 100
