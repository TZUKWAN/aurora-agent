"""Exporter for Dachuang application forms (DOCX and PDF)."""

import json
import logging
from typing import Dict

logger = logging.getLogger(__name__)


class DachuangDocxExporter:
    """Export Dachuang application to DOCX format."""

    def export(self, application: Dict, filepath: str) -> Dict:
        """Export application form to DOCX."""
        try:
            from aurora.business_plan.docx_builder import DocxBuilder
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt

            metadata = application.get("metadata", {})
            sections = application.get("sections", {})
            project_type = metadata.get("project_type", "innovation")

            doc = Document()

            # Setup styles
            style = doc.styles['Normal']
            font = style.font
            font.name = 'SimSun'
            font.size = Pt(12)

            # Cover
            for _ in range(4):
                doc.add_paragraph()

            title = doc.add_paragraph()
            run = title.add_run("大学生创新创业训练计划")
            run.bold = True
            run.font.size = Pt(22)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            subtitle = doc.add_paragraph()
            type_label = "创新训练项目" if project_type == "innovation" else "创业训练项目"
            run = subtitle.add_run(f"{type_label}申报书")
            run.bold = True
            run.font.size = Pt(20)
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph()

            # Basic info table
            info_data = [
                ["项目名称", metadata.get("project_name", "")],
                ["项目类型", type_label],
                ["负责人", metadata.get("leader", "")],
                ["指导教师", metadata.get("advisor", "")],
                ["所属院系", metadata.get("department", "")],
                ["申报等级", metadata.get("level", "校级")],
            ]
            table = doc.add_table(rows=len(info_data), cols=2)
            for i, (key, value) in enumerate(info_data):
                table.rows[i].cells[0].text = key
                table.rows[i].cells[1].text = str(value)

            doc.add_page_break()

            # Sections
            section_order = list(sections.keys())
            for sid in section_order:
                sec = sections.get(sid)
                if not sec:
                    continue
                content = sec.get("content", "") if isinstance(sec, dict) else str(sec)
                title_text = sec.get("title", sid) if isinstance(sec, dict) else sid
                if not content.strip():
                    continue

                heading = doc.add_heading(title_text, level=1)
                for run in heading.runs:
                    run.font.size = Pt(16)
                    run.font.bold = True

                for line in content.split("\n"):
                    if line.strip():
                        if line.startswith("【") and line.endswith("】"):
                            p = doc.add_heading(line, level=2)
                            p.runs[0].font.size = Pt(14)
                        elif line.startswith("- "):
                            doc.add_paragraph(line[2:], style='List Bullet')
                        elif line[0:1].isdigit() and "." in line[:3]:
                            doc.add_paragraph(line, style='List Number')
                        else:
                            doc.add_paragraph(line)

                doc.add_paragraph()

            doc.save(filepath)
            return {"success": True, "message": f"已导出到 {filepath}"}

        except ImportError:
            return {"success": False, "message": "python-docx is required"}
        except Exception as e:
            logger.exception("DOCX export failed")
            return {"success": False, "message": f"导出失败: {str(e)}"}


class DachuangPDFExporter:
    """Export Dachuang application to PDF using reportlab."""

    def export(self, application: Dict, filepath: str, theme: str = "minimal") -> Dict:
        """Export application form to PDF."""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import cm
            from reportlab.lib.colors import HexColor
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
            from reportlab.pdfbase import pdfmetrics

            # Get Chinese font
            chinese_font = "Helvetica"
            for name in ["SimSun", "SimHei", "Microsoft YaHei"]:
                try:
                    pdfmetrics.getFont(name)
                    chinese_font = name
                    break
                except Exception:
                    continue

            metadata = application.get("metadata", {})
            sections = application.get("sections", {})
            project_type = metadata.get("project_type", "innovation")

            doc = SimpleDocTemplate(filepath, pagesize=A4,
                                    leftMargin=2.5*cm, rightMargin=2.5*cm,
                                    topMargin=2.5*cm, bottomMargin=2.5*cm)

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'DachuangTitle',
                parent=styles['Heading1'],
                fontName=chinese_font,
                fontSize=18,
                textColor=HexColor("#333333"),
                spaceAfter=20,
            )
            heading_style = ParagraphStyle(
                'DachuangHeading',
                parent=styles['Heading2'],
                fontName=chinese_font,
                fontSize=14,
                textColor=HexColor("#2c5f8a"),
                spaceAfter=12,
            )
            body_style = ParagraphStyle(
                'DachuangBody',
                parent=styles['Normal'],
                fontName=chinese_font,
                fontSize=11,
                leading=18,
                spaceAfter=8,
            )

            elements = []

            # Title
            type_label = "创新训练项目" if project_type == "innovation" else "创业训练项目"
            elements.append(Paragraph("大学生创新创业训练计划", title_style))
            elements.append(Paragraph(f"{type_label}申报书", title_style))
            elements.append(Spacer(1, 1*cm))

            # Metadata
            info_lines = [
                f"项目名称：{metadata.get('project_name', '')}",
                f"负责人：{metadata.get('leader', '')}",
                f"指导教师：{metadata.get('advisor', '')}",
                f"所属院系：{metadata.get('department', '')}",
            ]
            for line in info_lines:
                elements.append(Paragraph(line, body_style))
            elements.append(Spacer(1, 0.5*cm))
            elements.append(PageBreak())

            # Sections
            for sid, sec in sections.items():
                content = sec.get("content", "") if isinstance(sec, dict) else str(sec)
                title_text = sec.get("title", sid) if isinstance(sec, dict) else sid
                if not content.strip():
                    continue

                elements.append(Paragraph(title_text, heading_style))
                elements.append(Spacer(1, 0.3*cm))

                for line in content.split("\n"):
                    if line.strip():
                        escaped = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                        elements.append(Paragraph(escaped, body_style))

                elements.append(Spacer(1, 0.5*cm))

            doc.build(elements)
            return {"success": True, "message": f"已导出到 {filepath}"}

        except ImportError:
            return {"success": False, "message": "reportlab is required"}
        except Exception as e:
            logger.exception("PDF export failed")
            return {"success": False, "message": f"导出失败: {str(e)}"}
