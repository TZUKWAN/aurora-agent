"""Enhanced DOCX builder for business plans."""

import logging
from typing import Dict, List, Optional

logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches, Cm, RGBColor
    from docx.enum.section import WD_ORIENT
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class DocxBuilder:
    """Build professional DOCX files from business plan data."""

    def __init__(self, theme: str = "default"):
        if not HAS_DOCX:
            raise ImportError("python-docx is required")
        self.theme = theme

    def build(self, plan: Dict, filepath: str, competition_name: str = "") -> Dict:
        """Build and save a DOCX file."""
        doc = Document()
        self._setup_styles(doc)

        metadata = plan.get("metadata", {})
        sections = plan.get("sections", {})

        # Cover page
        self._add_cover(doc, metadata, competition_name)

        # Table of contents placeholder
        self._add_toc(doc)

        # Sections
        section_order = [
            "executive_summary", "project_overview", "market_analysis",
            "product_service", "business_model", "marketing_strategy",
            "operation_plan", "team_introduction", "financial_analysis",
            "risk_assessment",
        ]

        for sid in section_order:
            sec = sections.get(sid)
            if not sec:
                continue
            content = sec.get("content", "") if isinstance(sec, dict) else str(sec)
            title = sec.get("title", sid) if isinstance(sec, dict) else sid
            if not content.strip():
                continue
            self._add_section(doc, title, content)

        doc.save(filepath)
        return {"success": True, "message": f"Saved to {filepath}"}

    def _setup_styles(self, doc):
        """Configure document styles."""
        style = doc.styles['Normal']
        font = style.font
        font.name = 'SimSun'
        font.size = Pt(11)

    def _add_cover(self, doc, metadata: Dict, competition_name: str):
        """Add cover page."""
        # Add some blank lines for vertical centering
        for _ in range(6):
            doc.add_paragraph()

        comp = competition_name or metadata.get("competition", "")
        title_text = f"{comp} 商业计划书" if comp else "商业计划书"

        title = doc.add_heading(title_text, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        date = metadata.get("generated_at", "")
        if date:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(date)
            run.font.size = Pt(14)

        track = metadata.get("track", "")
        if track:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"赛道: {track}")
            run.font.size = Pt(14)

        doc.add_page_break()

    def _add_toc(self, doc):
        """Add table of contents placeholder."""
        doc.add_heading("目录", level=1)
        doc.add_paragraph("（目录将在最终版本中自动生成）")
        doc.add_page_break()

    def _add_section(self, doc, title: str, content: str):
        """Add a section with title and formatted content."""
        doc.add_heading(title, level=1)

        # Split content by paragraphs
        paragraphs = content.split("\n")
        for para_text in paragraphs:
            para_text = para_text.strip()
            if not para_text:
                continue

            # Check if it's a sub-heading (starts with 【)
            if para_text.startswith("【") and para_text.endswith("】"):
                doc.add_heading(para_text.strip("【】"), level=2)
            elif para_text.startswith("【"):
                # Has sub-heading inline
                parts = para_text.split("】", 1)
                if len(parts) == 2:
                    doc.add_heading(parts[0].strip("【"), level=2)
                    rest = parts[1].strip()
                    if rest:
                        doc.add_paragraph(rest)
                else:
                    doc.add_paragraph(para_text)
            elif para_text.startswith("- ") or para_text.startswith("• "):
                # Bullet point
                doc.add_paragraph(para_text[2:], style='List Bullet')
            elif para_text.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")):
                # Numbered item
                doc.add_paragraph(para_text, style='List Number')
            else:
                doc.add_paragraph(para_text)

        doc.add_page_break()
