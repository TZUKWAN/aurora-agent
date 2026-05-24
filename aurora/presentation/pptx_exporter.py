"""Export PPTGenerator JSON structure to actual .pptx files."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

THEMES = {
    "business_blue": {
        "title_bg": RGBColor(0x1F, 0x3A, 0x5F),
        "title_fg": RGBColor(0xFF, 0xFF, 0xFF),
        "body_fg": RGBColor(0x33, 0x33, 0x33),
        "accent": RGBColor(0x2E, 0x86, 0xC1),
        "slide_bg": RGBColor(0xFF, 0xFF, 0xFF),
    },
    "tech": {
        "title_bg": RGBColor(0x1A, 0x1A, 0x2E),
        "title_fg": RGBColor(0x00, 0xFF, 0x88),
        "body_fg": RGBColor(0xE0, 0xE0, 0xE0),
        "accent": RGBColor(0x00, 0xD4, 0xAA),
        "slide_bg": RGBColor(0x16, 0x21, 0x3E),
    },
    "morandi": {
        "title_bg": RGBColor(0x8E, 0x8E, 0x9B),
        "title_fg": RGBColor(0xFF, 0xFF, 0xFF),
        "body_fg": RGBColor(0x4A, 0x4A, 0x4A),
        "accent": RGBColor(0xB5, 0x8B, 0x83),
        "slide_bg": RGBColor(0xF5, 0xF0, 0xEB),
    },
}


class PPTXExporter:
    """Convert PPTGenerator JSON output to .pptx files."""

    def __init__(self, theme: str = "business_blue"):
        self.theme_name = theme
        self.theme = THEMES.get(theme, THEMES["business_blue"])

    def export(self, ppt_content: dict, filepath: str) -> dict:
        """Export PPT JSON to .pptx file.

        Args:
            ppt_content: Output from PPTGenerator.generate()
            filepath: Output file path

        Returns:
            {"success": bool, "message": str}
        """
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        theme = self.theme

        for slide_data in ppt_content.get("slides", []):
            slide_id = slide_data.get("id", "")
            content = slide_data.get("content", [])

            if slide_id in ("cover", "closing"):
                self._build_title_slide(prs, slide_data, theme)
            else:
                self._build_content_slide(prs, slide_data, theme)

        prs.save(filepath)
        return {"success": True, "message": f"Exported to {filepath}"}

    def _build_title_slide(self, prs, slide_data, theme):
        """Build a title/cover/closing slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = theme["title_bg"]

        content = slide_data.get("content", [])
        top = Inches(2.0)

        for item in content:
            text = item.get("text", "")
            item_type = item.get("type", "bullet")

            if item_type in ("title",):
                self._add_textbox(slide, Inches(1), top, Inches(11.333), Inches(1.5),
                                  text, Pt(40), theme["title_fg"], PP_ALIGN.CENTER, bold=True)
                top += Inches(1.5)
            elif item_type in ("subtitle", "tagline"):
                self._add_textbox(slide, Inches(2), top, Inches(9.333), Inches(0.8),
                                  text, Pt(20), theme["title_fg"], PP_ALIGN.CENTER)
                top += Inches(0.8)
            else:
                self._add_textbox(slide, Inches(2), top, Inches(9.333), Inches(0.6),
                                  text, Pt(16), theme["title_fg"], PP_ALIGN.CENTER)
                top += Inches(0.6)

    def _build_content_slide(self, prs, slide_data, theme):
        """Build a content slide with title + bullet points."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = theme["slide_bg"]

        title_text = slide_data.get("title", "")
        # Title bar
        left_box = slide.shapes.add_shape(
            1, Inches(0), Inches(0), prs.slide_width, Inches(1.2)  # MSO_SHAPE.RECTANGLE
        )
        left_box.fill.solid()
        left_box.fill.fore_color.rgb = theme["title_bg"]
        left_box.line.fill.background()
        tf = left_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.color.rgb = theme["title_fg"]
        p.font.bold = True
        p.alignment = PP_ALIGN.LEFT
        tf.margin_left = Inches(0.5)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        # Bullet content
        content = slide_data.get("content", [])
        top = Inches(1.8)
        bullet_items = [item for item in content if item.get("type") != "title"]

        for item in bullet_items[:5]:
            text = item.get("text", "")
            item_type = item.get("type", "bullet")

            if item_type == "data":
                label = item.get("label", "")
                value = item.get("value", "")
                text = f"{label}: {value}"
            elif item_type == "timeline":
                phase = item.get("phase", "")
                desc = item.get("text", "")
                text = f"[{phase}] {desc}"

            self._add_textbox(slide, Inches(1), top, Inches(11), Inches(0.6),
                              text, Pt(18), theme["body_fg"], PP_ALIGN.LEFT)
            top += Inches(0.8)

    @staticmethod
    def _add_textbox(slide, left, top, width, height, text, font_size, color, alignment, bold=False):
        """Add a text box to a slide."""
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = font_size
        p.font.color.rgb = color
        p.font.bold = bold
        p.alignment = alignment
        return txBox
