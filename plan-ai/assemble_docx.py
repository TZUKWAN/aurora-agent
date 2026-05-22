#!/usr/bin/env python3
"""
批量商业计划书DOCX组装脚本
读取各项目的内容文件（支持.md/.json/.txt格式），生成格式规范的DOCX文件
"""

import os
import re
import json
import glob
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 项目清单
PROJECTS = {
    1: "智灵视界——基于少样本学习的工业精密部件缺陷零样本检测领航者",
    2: "孪生幻影——面向工业复杂场景的三维高斯溅射（3DGS）动态重建引擎",
    3: "极光智检——新能源电池极片涂布在线视觉多模态质检闭环软件",
    4: "具身智造——面向复杂工业环境的泛化视觉感知与空间推理大模型",
    5: "探微神算——基于时序图谱的工业设备剩余寿命预测与软测量平台",
    6: "储能先知——电化学储能电站热失控早期预警与动态寻优调度平台",
    7: "绿网调音师——虚拟电厂毫秒级源网荷储协同优化决策引擎",
    8: "元界智驾——面向自动驾驶Corner Case的端到端仿真测试生成引擎",
    9: "星轨卫士——低轨卫星星座通信链路拥塞智能预测与动态路由算法",
    10: "空域驭风——低空经济eVTOL多机协同防碰撞与航线动态规划大脑",
    11: "语境塑形——面向泛娱乐产业的剧本动态分支生成与逻辑一致性校验软件",
    12: "声境幻造——基于扩散模型的空间音频动态生成与个性化听觉渲染引擎",
    13: "幻视造物——面向3D资产生成的文本到高精度网格模型（Text-to-Mesh）转化系统",
    14: "芯流智核——面向大模型训练芯片的网络拓扑拥塞控制与仿真软件",
    15: "药界神农——基于量子化学计算的中药有效成分构效关系AI挖掘平台",
    16: "智测先锋——基于大语言模型与路径分析的代码级自动化测试用例生成中台",
    17: "云跃引擎——面向云原生架构的零配置智能部署与多维灰度发布平台",
    18: "影生万物——基于多模态大模型的跨平台营销短视频自动化生成引擎",
    19: "视界文枢——面向泛影视产业的AI剧本结构化拆解与动态分镜生成系统",
    20: "幻图智算——基于语义分割的泛电商视觉资产高通量自动化处理中台",
    21: "智绘经纬——基于结构化语言大模型的复杂商业架构多维可视化生成引擎",
    22: "知渊图谱——基于非结构化文本的领域实体关系抽取与动态知识图谱构建平台",
    23: "数语探微——面向商业决策的异构数据智能图表化解析与动态故事生成引擎",
    24: "商弈智境——基于多智能体强化学习的复杂商业环境动态演化与战略推演系统",
    25: "岁月留声——基于跨语种语音识别的非物质文化遗产口述史多模态数字建档平台",
    26: "元驾驭域——面向自动驾驶Corner Case的高保真数字孪生与极端场景自动生成系统",
    27: "星播智云——基于多模态情感计算的电商直播全链路智能运营与动态流控大脑",
    28: "心流源核——基于多模态生物反馈的职业倦怠干预与心理韧性数字疗法中台",
    29: "芯脉智连——基于强化学习的超大规模集成电路布线自动化寻优引擎",
    30: "智构视界——基于底层逻辑代码生成的超高密度商业框架图谱自动化渲染中台",
    31: "影刃识微——基于频域时空分析的AIGC深度伪造(Deepfake)防御引擎",
    32: "盲域星图——基于视觉语义分割的视障人群室内三维空间寻路语音中台",
    33: "银发智绘——基于大语言模型的老年阿尔茨海默症早期多模态语义衰退筛查引擎",
    34: "息流洞见——跨语种全球财经非结构化新闻的宏观经济指标前置预测大模型",
    35: "宠心译语——基于多模态声纹与行为图谱的跨物种情绪解码与交互大脑",
    36: "喵星食域——面向特异性肠道菌群的计算营养学自动配餐与动态演算软件",
    37: "烟火算盘——面向超微型商业体（地摊/夜市）的非结构化语音智能对账大脑",
    38: "质感寻源——基于材质纹理识别的闲置衣物数字化图谱与高阶胶囊衣橱生成引擎",
    39: "刺客雷达——基于商品历史曲线拟合与跨模态特征穿透的隐形价格陷阱扫描引擎",
    40: "闲光变现——基于海量长尾物品多模态估值的二手流转最优交易路径推演系统",
    41: "盲盒公交——基于城市全量公共交通时空图谱的低碳随机漫游(Citywalk)生成中台",
    42: "银发译林——面向老年人数字鸿沟的复杂操作界面实时语义降维与引导大模型",
    43: "智防深渊——面向老年群体的电信诈骗高维心理诱导反向侦破与熔断系统",
    44: "聚落寻根——基于多智能体演化的传统村落空间肌理自适应生成引擎",
    45: "群演矩阵——电影级战争场景百万智能体自主博弈与物理碰撞演算平台",
    46: "拆解先知——面向退役动力电池内部健康度(SOH)无损时序评估引擎",
    47: "冶金回响——电子微垃圾贵金属湿法回收工艺的AI多目标收率寻优模型",
    48: "数源织机——企业合成数据生成与隐私脱敏平台",
    49: "证链智核——电子证据整理与时间线生成系统",
}

OUTPUT_BASE = "D:/计划书AI/output"
FINAL_DIR = "D:/计划书AI/final"

CHAPTER_ORDER = [
    "第一章", "第二章", "第三章", "第四章", "第五章",
    "第六章", "第七章", "第八章", "第九章", "第十章"
]

CHAPTER_NAMES = [
    "企业基本情况", "项目介绍", "产品与服务", "市场分析", "竞争分析",
    "营销策略", "生产与运营", "财务分析", "风险分析与对策", "社会效益分析"
]


def find_content_file(project_id):
    """查找项目的内容文件"""
    patterns = [
        f"{OUTPUT_BASE}/batch*/project{project_id}/content.md",
        f"{OUTPUT_BASE}/batch*/project{project_id}/content.json",
        f"{OUTPUT_BASE}/batch*/project{project_id}/content.txt",
        f"{OUTPUT_BASE}/batch*/project{project_id}/business_plan_content.md",
        f"{OUTPUT_BASE}/batch*/project{project_id}_content.md",
        f"{OUTPUT_BASE}/batch*/project{project_id}_content.json",
        f"{OUTPUT_BASE}/batch*/project{project_id}_content.txt",
        f"{OUTPUT_BASE}/batch*/project{project_id}_chapters.md",
    ]
    for pattern in patterns:
        matches = glob.glob(pattern)
        if matches:
            return matches[0]

    # Check for chapter-separated files
    chapter_dir = f"{OUTPUT_BASE}/batch*/project{project_id}"
    for d in glob.glob(chapter_dir):
        chapter_files = sorted(glob.glob(os.path.join(d, "chapter*.txt")))
        if len(chapter_files) == 10:
            return ("chapters", chapter_files)

    return None


def parse_markdown_content(filepath):
    """解析markdown格式的内容"""
    with open(filepath, 'r', encoding='utf-8') as f:
        text = f.read()

    sections = []
    current_heading = ""
    current_content = []
    lines = text.split('\n')

    for line in lines:
        # Skip the top-level project title
        if line.startswith('# ') and not line.startswith('## '):
            continue
        if line.strip() == '---':
            continue

        # Detect headings
        if line.startswith('## ') or line.startswith('### '):
            if current_heading or current_content:
                content_text = '\n'.join(current_content).strip()
                if content_text:
                    sections.append({
                        'heading': current_heading,
                        'content': content_text,
                        'level': 2 if current_heading.startswith('## ') else 3
                    })
            current_heading = line.strip()
            current_content = []
        else:
            if line.strip():
                current_content.append(line.strip())

    # Don't forget the last section
    if current_heading or current_content:
        content_text = '\n'.join(current_content).strip()
        if content_text:
            sections.append({
                'heading': current_heading,
                'content': content_text,
                'level': 2 if current_heading.startswith('## ') else 3
            })

    return sections


def extract_content_text(content_val):
    """从content字段提取文本，兼容字符串和列表格式"""
    if isinstance(content_val, str):
        return content_val
    elif isinstance(content_val, list):
        texts = []
        for item in content_val:
            if isinstance(item, str):
                texts.append(item)
            elif isinstance(item, dict):
                texts.append(item.get('text', item.get('content', '')))
        return '\n'.join(texts)
    return str(content_val) if content_val else ''


def split_content_into_subsections(content_text):
    """将长文本按子标题拆分为子节"""
    if not content_text or len(content_text) < 50:
        return []

    subsections = []
    lines = content_text.split('\n')
    current_heading = ''
    current_content = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Detect sub-heading patterns
        is_sub = False
        if re.match(r'^[一二三四五六七八九十]+[、．.]', stripped):
            is_sub = True
        elif re.match(r'^\d+\.\d+\s', stripped):
            is_sub = True
        elif re.match(r'^[（(][一二三四五六七八九十]+[）)]', stripped):
            is_sub = True

        if is_sub:
            if current_content:
                subsections.append({
                    'heading': current_heading,
                    'content': ' '.join(current_content),
                    'level': 3
                })
            current_heading = stripped
            current_content = []
        else:
            current_content.append(stripped)

    if current_content:
        subsections.append({
            'heading': current_heading,
            'content': ' '.join(current_content),
            'level': 3
        })

    return subsections


def parse_json_content(filepath):
    """解析JSON格式的内容"""
    with open(filepath, 'r', encoding='utf-8') as f:
        data = json.load(f)

    sections = []

    if 'chapters' in data:
        chapters = data['chapters']
        # Handle dict format: {'chapter1': {'title': '...', 'content': '...'}, ...}
        if isinstance(chapters, dict):
            for ch_key in sorted(chapters.keys()):
                ch_data = chapters[ch_key]
                if isinstance(ch_data, dict):
                    ch_num_match = re.search(r'(\d+)', ch_key)
                    ch_num = int(ch_num_match.group(1)) if ch_num_match else 0
                    ch_title = f"第{num_to_chinese(ch_num)}章 {ch_data.get('title', '')}"
                    content_text = extract_content_text(ch_data.get('content', ''))
                    sections.append({
                        'heading': ch_title,
                        'content': content_text,
                        'level': 2
                    })
                    # Try to split content into sub-sections
                    sub_sections = split_content_into_subsections(content_text)
                    for sub in sub_sections:
                        sections.append(sub)
        # Handle list format: various structures
        elif isinstance(chapters, list):
            for i_ch, chapter in enumerate(chapters, 1):
                if isinstance(chapter, str):
                    continue
                if not isinstance(chapter, dict):
                    continue

                ch_num = chapter.get('chapter', 0)
                if ch_num:
                    ch_title = f"第{num_to_chinese(ch_num)}章 {chapter.get('title', '')}"
                else:
                    title = chapter.get('title', '')
                    ch_title = f"第{num_to_chinese(i_ch)}章 {title}" if title else f"第{num_to_chinese(i_ch)}章"

                # Check if chapters have 'sections' array or direct 'content' string
                if 'sections' in chapter:
                    sections.append({
                        'heading': ch_title,
                        'content': '',
                        'level': 2
                    })
                    for sec in chapter['sections']:
                        if isinstance(sec, dict):
                            content_text = extract_content_text(sec.get('content', ''))
                            sections.append({
                                'heading': sec.get('heading', ''),
                                'content': content_text,
                                'level': 3
                            })
                elif 'content' in chapter:
                    content_text = extract_content_text(chapter['content'])
                    sections.append({
                        'heading': ch_title,
                        'content': content_text,
                        'level': 2
                    })
                    sub_sections = split_content_into_subsections(content_text)
                    for sub in sub_sections:
                        sections.append(sub)
                else:
                    sections.append({
                        'heading': ch_title,
                        'content': '',
                        'level': 2
                    })
    elif 'sections' in data:
        for sec in data['sections']:
            heading = sec.get('heading', '')
            level_val = sec.get('level', 3)
            if level_val == 1 or ('第' in heading and '章' in heading):
                level = 2
            else:
                level = 3
            content_text = extract_content_text(sec.get('content', sec.get('text', '')))
            sections.append({
                'heading': heading,
                'content': content_text,
                'level': level
            })

    return sections


def parse_txt_content(filepath):
    """解析纯文本格式的内容"""
    with open(filepath, 'r', encoding='utf-8') as f:
        text = f.read()

    sections = []
    lines = text.split('\n')
    current_heading = ""
    current_content = []

    for line in lines:
        # Detect chapter headings
        is_heading = False
        level = 3

        if re.match(r'^第[一二三四五六七八九十]+章', line.strip()):
            is_heading = True
            level = 2
        elif re.match(r'^[一二三四五六七八九十]+[、.]', line.strip()):
            is_heading = True
            level = 3
        elif re.match(r'^\d+\.\d+\s', line.strip()):
            is_heading = True
            level = 3
        # Short lines that look like headings
        elif len(line.strip()) < 30 and line.strip() and not line.strip().startswith('【') and current_content:
            # Check if it's a section title pattern
            pass

        if is_heading:
            if current_heading or current_content:
                content_text = '\n'.join(current_content).strip()
                if content_text:
                    sections.append({
                        'heading': current_heading,
                        'content': content_text,
                        'level': level
                    })
            current_heading = line.strip()
            current_content = []
        else:
            if line.strip():
                current_content.append(line.strip())

    if current_heading or current_content:
        content_text = '\n'.join(current_content).strip()
        if content_text:
            sections.append({
                'heading': current_heading,
                'content': content_text,
                'level': 3
            })

    return sections


def parse_chapter_files(chapter_files):
    """解析分章节的txt文件"""
    sections = []
    for cf in chapter_files:
        with open(cf, 'r', encoding='utf-8') as f:
            text = f.read()

        basename = os.path.basename(cf)
        ch_num = int(re.search(r'chapter(\d+)', basename).group(1))
        ch_title = f"第{num_to_chinese(ch_num)}章 {CHAPTER_NAMES[ch_num - 1]}"

        sections.append({
            'heading': ch_title,
            'content': '',
            'level': 2
        })

        lines = text.strip().split('\n')
        current_sub = ""
        current_para = []

        for line in lines:
            stripped = line.strip()
            if not stripped:
                if current_para:
                    content_text = ' '.join(current_para)
                    if current_sub:
                        sections.append({
                            'heading': current_sub,
                            'content': content_text,
                            'level': 3
                        })
                    else:
                        sections[-1]['content'] = content_text if not sections[-1]['content'] else sections[-1]['content'] + '\n\n' + content_text
                    current_para = []
                continue

            # Detect sub-headings
            if re.match(r'^[一二三四五六七八九十]+[、.]', stripped) or re.match(r'^\d+\.\d+', stripped):
                if current_para and current_sub:
                    sections.append({
                        'heading': current_sub,
                        'content': ' '.join(current_para),
                        'level': 3
                    })
                    current_para = []
                current_sub = stripped
            else:
                current_para.append(stripped)

        if current_para and current_sub:
            sections.append({
                'heading': current_sub,
                'content': ' '.join(current_para),
                'level': 3
            })
        elif current_para:
            content_text = ' '.join(current_para)
            if sections[-1]['content']:
                sections[-1]['content'] += '\n\n' + content_text
            else:
                sections[-1]['content'] = content_text

    return sections


def num_to_chinese(n):
    """数字转中文"""
    mapping = {1: '一', 2: '二', 3: '三', 4: '四', 5: '五',
               6: '六', 7: '七', 8: '八', 9: '九', 10: '十'}
    return mapping.get(n, str(n))


def set_cell_font(cell, text, font_name='宋体', font_size=10.5, bold=False):
    """设置表格单元格字体"""
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def add_page_number(doc):
    """添加页码"""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run()
        fld_char1 = OxmlElement('w:fldChar')
        fld_char1.set(qn('w:fldCharType'), 'begin')
        run._element.append(fld_char1)

        run2 = p.add_run()
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'PAGE'
        run2._element.append(instr)

        run3 = p.add_run()
        fld_char2 = OxmlElement('w:fldChar')
        fld_char2.set(qn('w:fldCharType'), 'end')
        run3._element.append(fld_char2)


def create_docx(project_id, project_name, sections, output_path):
    """生成DOCX文件"""
    doc = Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # Set paragraph format
    pf = style.paragraph_format
    pf.line_spacing = Pt(22)
    pf.space_after = Pt(6)
    pf.first_line_indent = Cm(0.74)

    # Cover page
    for _ in range(6):
        doc.add_paragraph('')

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(project_name)
    title_run.font.name = '黑体'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run('湖北省大学生创业扶持项目商业计划书')
    subtitle_run.font.name = '黑体'
    subtitle_run.font.size = Pt(18)
    subtitle_run.font.bold = True
    subtitle_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    for _ in range(4):
        doc.add_paragraph('')

    info_items = [
        f'项目名称：{project_name}',
        '申报单位：【待填写】',
        '法定代表人：【待填写】',
        '申报日期：2026年4月',
    ]
    for item in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(item)
        run.font.name = '仿宋'
        run.font.size = Pt(14)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

    # Page break after cover
    doc.add_page_break()

    # Table of contents
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_run = toc_title.add_run('目  录')
    toc_run.font.name = '黑体'
    toc_run.font.size = Pt(16)
    toc_run.font.bold = True
    toc_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    for i, (ch, name) in enumerate(zip(CHAPTER_ORDER, CHAPTER_NAMES), 1):
        p = doc.add_paragraph()
        run = p.add_run(f'{ch}  {name}')
        run.font.name = '宋体'
        run.font.size = Pt(12)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    doc.add_page_break()

    # Add content sections
    for sec in sections:
        heading_text = sec['heading'].lstrip('#').strip()
        content_text = sec['content']

        if not heading_text and not content_text:
            continue

        # Add heading
        if heading_text:
            clean_heading = re.sub(r'^#+\s*', '', heading_text)
            if sec['level'] == 2:
                h = doc.add_heading(clean_heading, level=1)
                for run in h.runs:
                    run.font.name = '黑体'
                    run.font.size = Pt(16)
                    run.font.bold = True
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            else:
                h = doc.add_heading(clean_heading, level=2)
                for run in h.runs:
                    run.font.name = '黑体'
                    run.font.size = Pt(14)
                    run.font.bold = True
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

        # Add content paragraphs
        if content_text:
            paragraphs = content_text.split('\n')
            for para_text in paragraphs:
                para_text = para_text.strip()
                if not para_text:
                    continue

                p = doc.add_paragraph()
                # Handle first line indent for body text
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0.74)
                pf.line_spacing = Pt(22)

                run = p.add_run(para_text)
                run.font.name = '宋体'
                run.font.size = Pt(12)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # Add page numbers
    add_page_number(doc)

    # Save
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def main():
    os.makedirs(FINAL_DIR, exist_ok=True)

    success = 0
    failed = 0
    skipped = 0

    for project_id in range(1, 50):
        project_name = PROJECTS.get(project_id, f"项目{project_id}")
        print(f"\n[{project_id}/49] 处理: {project_name[:40]}...")

        content_file = find_content_file(project_id)

        if content_file is None:
            print(f"  ⚠ 跳过 - 未找到内容文件")
            skipped += 1
            continue

        try:
            # Parse content based on format
            if isinstance(content_file, tuple) and content_file[0] == "chapters":
                sections = parse_chapter_files(content_file[1])
            elif content_file.endswith('.json'):
                sections = parse_json_content(content_file)
            elif content_file.endswith('.md'):
                sections = parse_markdown_content(content_file)
            else:
                sections = parse_txt_content(content_file)

            if not sections:
                print(f"  ⚠ 跳过 - 解析后无内容")
                skipped += 1
                continue

            # Generate output filename
            safe_name = re.sub(r'[\\/:*?"<>|]', '', project_name)
            output_filename = f"【{safe_name}】湖北省大学生创业扶持项目商业计划书.docx"
            output_path = os.path.join(FINAL_DIR, output_filename)

            create_docx(project_id, project_name, sections, output_path)

            file_size = os.path.getsize(output_path)
            print(f"  ✓ 成功 - {output_filename} ({file_size/1024:.1f}KB)")
            success += 1

        except Exception as e:
            print(f"  ✗ 失败 - {str(e)}")
            failed += 1

    print(f"\n{'='*60}")
    print(f"处理完成: 成功 {success}, 失败 {failed}, 跳过 {skipped}")
    print(f"输出目录: {FINAL_DIR}")


if __name__ == '__main__':
    main()
