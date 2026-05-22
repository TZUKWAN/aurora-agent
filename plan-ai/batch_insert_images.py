#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
batch_insert_images.py
用 visual_chart_engine + screenshot.py 为49个项目生成视觉上互不相同的图片并插入DOCX。
不再依赖PPT引擎，使用 matplotlib + Pillow 直接绘制，速度快、样式多样。
"""

import json, os, re, sys, time
from pathlib import Path
from PIL import Image as PILImage
import openpyxl
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---- Paths ----
BASE_DIR = Path(r"D:\计划书AI")
OUTPUT_DIR = BASE_DIR / "output_v2"
IMAGE_BASE = OUTPUT_DIR / "images"
XLSX_PATH = BASE_DIR / "省补贴项目清单.xlsx"
PROGRESS_FILE = BASE_DIR / "img_insert_progress.json"

SKILL_SS = Path(r"C:\Users\lauze\.claude\skills\business-plan-writer\subskills\bpw-product-screenshots\scripts")
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(SKILL_SS))

from visual_chart_engine import generate_figure, FIGURE_GENERATORS
from screenshot import generate_set

# 截图映射: (fig_key, screenshot_index)
SCREENSHOT_FIGS = [
    ("图2-7", 0), ("图2-8", 1), ("图2-9", 2), ("图2-10", 3),
    ("图2-14", 4),
] + [(f"图2-{i}", i - 22 + 5) for i in range(22, 52)]


def load_projects():
    wb = openpyxl.load_workbook(str(XLSX_PATH))
    ws = wb.active
    projects = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        title = r[0] or ''
        name = title.split('——')[0].strip() if '——' in title else title[:20]
        company = r[2] or f'{name}科技有限公司'
        projects.append({'name': name, 'title': title, 'desc': r[1] or '',
                         'product_name': name, 'company_name': company})
    return projects


def find_docx(name):
    for f in os.listdir(str(OUTPUT_DIR)):
        if f.endswith('.docx') and not f.startswith('~') and name in f:
            return OUTPUT_DIR / f
    return None


def process_project(idx, project):
    """处理单个项目：生成所有图片 + 插入DOCX"""
    name = project['name']
    img_dir = IMAGE_BASE / f"p{idx}"
    img_dir.mkdir(parents=True, exist_ok=True)

    # ---- 1. 生成38种图表 ----
    chart_map = {}
    for fig_key in FIGURE_GENERATORS:
        out_path = img_dir / f"{fig_key.replace('-', '_')}.png"
        if not out_path.exists():
            generate_figure(name, fig_key, str(out_path))
        chart_map[fig_key] = str(out_path)

    # ---- 2. 生成35张产品截图 ----
    ss_dir = img_dir / "screenshots"
    manifest_path = ss_dir / "manifest.json"
    if manifest_path.exists():
        try:
            screenshots = json.loads(manifest_path.read_text(encoding='utf-8'))
            if len(screenshots) < 35:
                raise ValueError("insufficient")
        except:
            screenshots = generate_set(
                {'product_name': project['product_name'],
                 'company_name': project['company_name']}, ss_dir, count=35)
            manifest_path.write_text(json.dumps(screenshots, ensure_ascii=False, indent=2), encoding='utf-8')
    else:
        screenshots = generate_set(
            {'product_name': project['product_name'],
             'company_name': project['company_name']}, ss_dir, count=35)
        manifest_path.write_text(json.dumps(screenshots, ensure_ascii=False, indent=2), encoding='utf-8')

    # ---- 3. 构建完整映射 ----
    image_map = dict(chart_map)
    for fig_key, sidx in SCREENSHOT_FIGS:
        if sidx < len(screenshots):
            image_map[fig_key] = screenshots[sidx]['path']

    # ---- 4. 插入DOCX ----
    docx_path = find_docx(name)
    if not docx_path:
        return 0, 0

    doc = Document(str(docx_path))
    inserted = 0
    paras = doc.paragraphs

    for i, p in enumerate(paras):
        m = re.match(r'^(图\d+-\d+)\s', p.text.strip())
        if not m:
            continue
        fig_key = m.group(1)
        img_path = image_map.get(fig_key)
        if not img_path or not os.path.exists(img_path):
            continue

        # 图注前的空段落作为图片位置
        if i > 0 and not paras[i - 1].text.strip():
            target = paras[i - 1]
            # 清除残留
            for child in list(target._element):
                tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                if tag == 'r':
                    target._element.remove(child)

            # 图片尺寸：竖屏小，横屏大
            try:
                w, h = PILImage.open(img_path).size
                width = Inches(2.8) if h > w * 1.5 else Inches(5.2)
            except:
                width = Inches(5.2)

            run = target.add_run()
            try:
                run.add_picture(img_path, width=width)
                inserted += 1
            except Exception as e:
                print(f"    insert {fig_key} fail: {e}")

    # 居中图片段落
    for p in doc.paragraphs:
        if not p.text.strip():
            has_img = any(r._element.findall(qn('w:drawing')) for r in p.runs)
            if has_img:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(docx_path))
    return len(image_map), inserted


def main():
    print("=" * 60)
    print(" 湖北省创业扶持项目 - 图片生成与插入 (visual_chart_engine)")
    print("=" * 60)

    progress = {}
    if PROGRESS_FILE.exists():
        try:
            progress = json.loads(PROGRESS_FILE.read_text(encoding='utf-8'))
        except:
            pass

    projects = load_projects()
    total = len(projects)
    done = sum(1 for v in progress.values() if v.get('done'))
    print(f"共 {total} 个项目, 已完成 {done}, 剩余 {total - done}")

    t_start = time.time()

    for idx in range(total):
        pid = f"p{idx}"
        if progress.get(pid, {}).get('done'):
            continue

        t0 = time.time()
        try:
            total_imgs, inserted = process_project(idx, projects[idx])
        except Exception as e:
            print(f"[{idx+1}/{total}] {projects[idx]['name']} FAILED: {e}")
            continue

        elapsed = time.time() - t0
        progress[pid] = {'done': True, 'imgs': total_imgs, 'inserted': inserted, 'time': round(elapsed, 1)}
        PROGRESS_FILE.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding='utf-8')

        done += 1
        eta = (time.time() - t_start) / done * (total - done) / 60
        print(f"[{done}/{total}] {projects[idx]['name']}: {total_imgs}张生成, {inserted}张插入, "
              f"{elapsed:.0f}s, ETA {eta:.1f}min")

    elapsed_total = time.time() - t_start
    total_inserted = sum(v.get('inserted', 0) for v in progress.values())
    print(f"\n{'='*60}")
    print(f"完成: {done}/{total} 项目, 共插入 {total_inserted} 张图片, 耗时 {elapsed_total/60:.1f}分钟")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
