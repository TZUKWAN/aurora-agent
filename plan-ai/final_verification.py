#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Final verification of all 49 DOCX files."""
import re
import zipfile
from pathlib import Path
from docx import Document

OUTPUT_DIR = Path(r'D:\计划书AI\output_v2')
docx_files = [f for f in OUTPUT_DIR.glob('*.docx') if not f.name.startswith('~$')]

issues = []

for docx_path in docx_files:
    doc = Document(str(docx_path))
    
    # 1. Count figure captions
    fig_captions = [p.text.strip() for p in doc.paragraphs if re.match(r'^图\d+-\d+\s', p.text.strip())]
    
    # 2. Count images in body paragraphs
    img_count = sum(1 for p in doc.paragraphs if any(
        r._element.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing') for r in p.runs))
    
    # 3. Count tables
    table_count = len(doc.tables)
    
    # 4. Check for remaining placeholders
    placeholder_count = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if '【图' in text or '【表' in text or '请插入图片' in text or '请插入表格' in text:
            placeholder_count += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text.strip()
                    if '【图' in text or '【表' in text or '请插入图片' in text or '请插入表格' in text:
                        placeholder_count += 1
    
    # 5. Check for orphaned images in zip
    with zipfile.ZipFile(docx_path, 'r') as z:
        doc_xml = z.read('word/document.xml').decode('utf-8')
        rels = z.read('word/_rels/document.xml.rels').decode('utf-8')
        
        embeds = set(re.findall(r'embed="(.*?)"', doc_xml))
        rel_map = {}
        for match in re.findall(r'Id="(.*?)"[^>]*Type=".*?image"[^>]*Target="media/(.*?)"', rels):
            rel_map[match[0]] = match[1]
        
        referenced = set(rel_map.get(rid) for rid in embeds if rid in rel_map)
        all_media = set(f.split('/')[-1] for f in z.namelist() if f.startswith('word/media/'))
        orphaned = all_media - referenced
    
    # Check for issues
    doc_issues = []
    if len(fig_captions) != 33:
        doc_issues.append(f'图注数量异常: {len(fig_captions)} (应为33)')
    if img_count != 33:
        doc_issues.append(f'图片数量异常: {img_count} (应为33)')
    if table_count != 3:
        doc_issues.append(f'表格数量异常: {table_count} (应为3)')
    if placeholder_count > 0:
        doc_issues.append(f'剩余占位符: {placeholder_count}')
    if len(orphaned) > 0:
        doc_issues.append(f'Orphaned图片: {len(orphaned)}')
    
    if doc_issues:
        issues.append((docx_path.name[:30], doc_issues))

if issues:
    print(f'发现 {len(issues)} 个文档存在问题:')
    for name, doc_issues in issues:
        print(f'  {name}...')
        for issue in doc_issues:
            print(f'    - {issue}')
else:
    print(f'全部 {len(docx_files)} 个文档验证通过!')
    print(f'  - 每个文档包含 33 张图 + 33 个图注')
    print(f'  - 每个文档包含 3 个真实表格')
    print(f'  - 无剩余占位符')
    print(f'  - 无 orphaned 图片')
