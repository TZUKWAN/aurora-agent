#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Remove placeholder tables from all DOCX files.
A placeholder table is defined as a table whose only text contains
bracketed placeholders like 【图 X-X 区域 - 请插入图片】
"""
import os
import re
import time
from pathlib import Path
from docx import Document

BASE_DIR = Path(r"D:\计划书AI")
OUTPUT_DIR = BASE_DIR / "output_v2"


def is_placeholder_table(table):
    """Check if a table contains only placeholder text."""
    all_texts = []
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                all_texts.append(para.text.strip())
    
    full_text = ' '.join(all_texts)
    
    # If empty, it's a placeholder
    if not full_text.strip():
        return True
    
    # If it contains bracketed placeholders and no real content
    has_placeholder = '【图' in full_text or '【表' in full_text or '请插入' in full_text
    
    # Check for real content (text without brackets and not empty)
    real_content = [t for t in all_texts if t and '【' not in t and '请插入' not in t]
    
    return has_placeholder and len(real_content) == 0


def remove_table(doc, table):
    """Remove a table from the document."""
    table._element.getparent().remove(table._element)


def clean_docx(docx_path):
    """Remove all placeholder tables from a DOCX file."""
    doc = Document(str(docx_path))
    
    removed = 0
    # Iterate in reverse to avoid index shifting issues
    for table in reversed(doc.tables):
        if is_placeholder_table(table):
            remove_table(doc, table)
            removed += 1
    
    doc.save(str(docx_path))
    return removed


def main():
    print("=" * 60)
    print(" 删除占位符表格")
    print("=" * 60)
    
    docx_files = [f for f in OUTPUT_DIR.glob('*.docx') if not f.name.startswith('~$')]
    total = len(docx_files)
    total_removed = 0
    
    t_start = time.time()
    for i, docx_path in enumerate(docx_files):
        t0 = time.time()
        removed = clean_docx(docx_path)
        elapsed = time.time() - t0
        total_removed += removed
        print(f"[{i+1}/{total}] {docx_path.name[:20]}...: 删除 {removed} 个占位符表格 ({elapsed:.1f}s)")
    
    total_elapsed = time.time() - t_start
    print(f"\n{'='*60}")
    print(f"完成: {total}个文档, 共删除 {total_removed} 个占位符表格")
    print(f"耗时 {total_elapsed:.1f}s")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
