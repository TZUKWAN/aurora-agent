#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_images_and_insert.py
为49个湖北省大学生创业扶持项目批量生成图片并插入DOCX文件。

阶段1: 生成产品截图 (Playwright) - 每项目35张
阶段2: 生成学术图表 (PowerPoint COM) - 每项目38张
阶段3: 插入图片到DOCX (python-docx) - 每项目73张图
"""

import json, os, re, sys, time, random
from pathlib import Path
from PIL import Image as PILImage

# ---- Paths ----
BASE_DIR = Path(r"D:\计划书AI")
OUTPUT_DIR = BASE_DIR / "output_v2"
IMAGE_BASE = OUTPUT_DIR / "images"
XLSX_PATH = BASE_DIR / "省补贴项目清单.xlsx"
PROGRESS_FILE = BASE_DIR / "image_gen_progress.json"

SKILL_BASE = Path(r"C:\Users\lauze\.claude\skills\business-plan-writer")
SCREENSHOT_DIR = SKILL_BASE / "subskills" / "bpw-product-screenshots" / "scripts"
CHART_DIR = SKILL_BASE / "scripts"

sys.path.insert(0, str(SCREENSHOT_DIR))
sys.path.insert(0, str(CHART_DIR))

# ---- 图表类图片映射: figure_key -> (chart_type, title_suffix) ----
CHART_FIGURES = {
    "图1-1": ("architecture", "公司组织架构"),
    "图1-2": ("gantt", "发展战略路线"),
    "图1-3": ("swot", None),
    "图1-4": ("competition", "波特六力模型"),
    "图1-5": ("gantt", "研发路线"),
    "图1-6": ("value_chain", "产品矩阵"),
    "图1-7": ("architecture", "服务体系架构"),
    "图1-8": ("value_chain", "营销渠道"),
    "图1-9": ("bmc", None),
    "图2-1": ("value_chain", "产业链"),
    "图2-2": ("competition", "市场规模趋势"),
    "图2-3": ("gantt", "政策支持时间线"),
    "图2-4": ("architecture", "产品功能架构"),
    "图2-5": ("architecture", "系统技术架构"),
    "图2-6": ("value_chain", "核心算法流程"),
    "图2-11": ("competition", "目标客户画像"),
    "图2-12": ("value_chain", "客户需求分析"),
    "图2-13": ("competition", "产品优势对比"),
    "图2-15": ("competition", "应用效果对比"),
    "图2-16": ("competition", "市场占有率"),
    "图2-17": ("competition", "行业竞争格局"),
    "图2-18": ("bmc", "收入构成"),
    "图2-19": ("competition", "收入增长趋势"),
    "图2-20": ("competition", "竞争格局矩阵"),
    "图2-21": ("competition", "竞争对手对标"),
    "图2-52": ("architecture", "系统整体架构"),
    "图2-53": ("architecture", "数据流架构"),
    "图2-54": ("value_chain", "核心算法架构"),
    "图2-55": ("architecture", "微服务架构"),
    "图2-56": ("architecture", "数据库架构"),
    "图2-57": ("architecture", "安全架构"),
    "图2-58": ("architecture", "部署架构"),
    "图2-59": ("architecture", "API架构"),
    "图2-60": ("architecture", "前端架构"),
    "图2-61": ("architecture", "运维监控架构"),
    "图3-1": ("bmc", "资金用途分布"),
    "图3-2": ("gantt", "融资里程碑"),
    "图4-1": ("swot", None),
}

# ---- 截图类图片映射: (figure_key, screenshot_manifest_index) ----
SCREENSHOT_FIGURES = [
    ("图2-7", 0), ("图2-8", 1), ("图2-9", 2), ("图2-10", 3),
    ("图2-14", 4),
] + [(f"图2-{i}", i - 22 + 5) for i in range(22, 52)]


def load_projects():
    import openpyxl
    wb = openpyxl.load_workbook(str(XLSX_PATH))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    projects = []
    for r in rows:
        title = r[0] or ''
        desc = r[1] or ''
        company = r[2] or ''
        name = title.split('——')[0].strip() if '——' in title else title[:20]
        projects.append({
            'title': title, 'name': name, 'desc': desc,
            'company': company,
            'product_name': name,
            'company_name': company or f'{name}科技有限公司',
        })
    return projects


def find_docx_for_project(name):
    for f in os.listdir(str(OUTPUT_DIR)):
        if f.endswith('.docx') and not f.startswith('~') and name in f:
            return OUTPUT_DIR / f
    return None


# ---- 图表数据生成器 ----
def gen_chart_data(project, chart_type, title_suffix):
    n = project['name']

    if chart_type == "swot":
        return {
            "s": f"{n}技术团队实力雄厚\n产品创新性强\n市场先发优势明显",
            "w": "品牌知名度有待提升\n资金规模相对有限\n渠道网络尚待完善",
            "o": "政策大力支持产业发展\n市场需求呈现快速增长\n技术窗口期带来机遇",
            "t": "行业巨头可能进入\n技术更新迭代速度快\n客户获取成本持续上升",
        }

    if chart_type == "gantt":
        phase_sets = {
            "发展战略路线": [
                {"name": "产品研发阶段", "start_month": 1, "duration_months": 8, "milestones": ["MVP发布"]},
                {"name": "市场验证阶段", "start_month": 6, "duration_months": 6, "milestones": ["首批客户"]},
                {"name": "规模扩张阶段", "start_month": 12, "duration_months": 8, "milestones": ["用户突破10万"]},
                {"name": "生态构建阶段", "start_month": 20, "duration_months": 8, "milestones": ["行业标杆"]},
            ],
            "研发路线": [
                {"name": "核心技术攻关", "start_month": 1, "duration_months": 6},
                {"name": "原型系统开发", "start_month": 4, "duration_months": 4},
                {"name": "内测与优化", "start_month": 7, "duration_months": 3},
                {"name": "公测与迭代", "start_month": 9, "duration_months": 4},
                {"name": "正式发布", "start_month": 12, "duration_months": 2},
            ],
            "政策支持时间线": [
                {"name": "政策研究立项", "start_month": 1, "duration_months": 3},
                {"name": "政策草案制定", "start_month": 3, "duration_months": 4},
                {"name": "征求意见阶段", "start_month": 6, "duration_months": 2},
                {"name": "政策正式发布", "start_month": 8, "duration_months": 1},
                {"name": "落地实施推广", "start_month": 9, "duration_months": 6},
            ],
            "融资里程碑": [
                {"name": "种子轮融资", "start_month": 1, "duration_months": 3},
                {"name": "产品研发投入", "start_month": 3, "duration_months": 6},
                {"name": "天使轮融资", "start_month": 8, "duration_months": 3},
                {"name": "市场拓展投入", "start_month": 10, "duration_months": 6},
                {"name": "A轮融资准备", "start_month": 16, "duration_months": 4},
            ],
        }
        return {"phases": phase_sets.get(title_suffix, phase_sets["发展战略路线"])}

    if chart_type == "bmc":
        return {
            "key_partners": "云服务商、数据供应商、行业合作伙伴",
            "key_activities": "核心技术研发、产品迭代、客户服务",
            "key_resources": "技术专利、核心团队、用户数据",
            "value_propositions": f"{n}提供智能化解决方案，提升效率降低成本",
            "customer_relationships": "技术支持、定制服务、社区运营",
            "channels": "直销团队、线上推广、合作伙伴",
            "customer_segments": "中大型企业、政府机构、产业园区",
            "cost_structure": "研发投入50%、市场推广25%、运营成本25%",
            "revenue_streams": "SaaS订阅60%、定制开发25%、数据服务15%",
        }

    if chart_type == "value_chain":
        node_sets = {
            "产业链": [{"name": "上游供应商", "type": "supplier"}, {"name": "核心技术", "type": "core"},
                       {"name": n, "type": "center"}, {"name": "下游客户", "type": "customer"}, {"name": "终端用户", "type": "end_user"}],
            "产品矩阵": [{"name": "基础版", "type": "product"}, {"name": "专业版", "type": "product"},
                         {"name": f"{n}平台", "type": "center"}, {"name": "企业版", "type": "product"}, {"name": "定制版", "type": "product"}],
            "营销渠道": [{"name": "线上推广", "type": "channel"}, {"name": "直销团队", "type": "channel"},
                         {"name": n, "type": "center"}, {"name": "合作伙伴", "type": "channel"}, {"name": "行业展会", "type": "channel"}],
            "核心算法流程": [{"name": "数据采集", "type": "process"}, {"name": "数据预处理", "type": "process"},
                            {"name": "核心算法", "type": "center"}, {"name": "结果输出", "type": "process"}, {"name": "反馈优化", "type": "process"}],
            "客户需求分析": [{"name": "需求调研", "type": "input"}, {"name": "需求分析", "type": "process"},
                            {"name": "解决方案", "type": "center"}, {"name": "实施交付", "type": "output"}, {"name": "持续服务", "type": "service"}],
            "核心算法架构": [{"name": "输入层", "type": "layer"}, {"name": "特征提取", "type": "layer"},
                            {"name": "核心模型", "type": "center"}, {"name": "推理引擎", "type": "layer"}, {"name": "输出层", "type": "layer"}],
        }
        nodes = node_sets.get(title_suffix, node_sets["产业链"])
        edges = [{"from": i, "to": i + 1} for i in range(len(nodes) - 1)]
        return {"nodes": nodes, "edges": edges}

    if chart_type == "competition":
        player_sets = {
            "波特六力模型": [
                {"name": "现有竞争者", "market_share": 30, "innovation": 60},
                {"name": "潜在进入者", "market_share": 10, "innovation": 70},
                {"name": "替代品威胁", "market_share": 15, "innovation": 50},
                {"name": "供应商议价", "market_share": 20, "innovation": 40},
                {"name": "买方议价", "market_share": 15, "innovation": 30},
                {"name": n, "market_share": 10, "innovation": 85},
            ],
            "市场规模趋势": [
                {"name": "2024年", "market_share": 40, "innovation": 50},
                {"name": "2025年", "market_share": 55, "innovation": 60},
                {"name": "2026年", "market_share": 70, "innovation": 70},
                {"name": "2027年", "market_share": 85, "innovation": 75},
                {"name": "2028年", "market_share": 100, "innovation": 80},
            ],
            "目标客户画像": [
                {"name": "大型企业", "market_share": 35, "innovation": 80},
                {"name": "中型企业", "market_share": 30, "innovation": 65},
                {"name": "小型企业", "market_share": 20, "innovation": 50},
                {"name": "政府机构", "market_share": 10, "innovation": 70},
                {"name": "其他", "market_share": 5, "innovation": 40},
            ],
        }
        default_players = [
            {"name": n, "market_share": random.randint(10, 25), "innovation": random.randint(80, 95)},
            {"name": "竞品A", "market_share": random.randint(20, 35), "innovation": random.randint(55, 75)},
            {"name": "竞品B", "market_share": random.randint(15, 25), "innovation": random.randint(50, 65)},
            {"name": "竞品C", "market_share": random.randint(10, 20), "innovation": random.randint(40, 55)},
            {"name": "传统方案", "market_share": random.randint(5, 15), "innovation": random.randint(25, 45)},
        ]
        return {"players": player_sets.get(title_suffix, default_players)}

    return {}


# ---- 截图生成 ----
def gen_screenshots(project, img_dir, count=35):
    from screenshot import generate_set
    ss_dir = img_dir / "screenshots"
    manifest_path = ss_dir / "manifest.json"

    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            if len(manifest) >= count:
                return manifest
        except:
            pass

    project_info = {
        "product_name": project["product_name"],
        "company_name": project["company_name"],
    }
    results = generate_set(project_info, ss_dir, count=count)
    manifest_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  Screenshots: {len(results)} generated")
    return results


# ---- 图表生成 ----
def gen_charts(engine, project, img_dir):
    from ppt_diagram_engine import PPTDiagramEngine
    name = project["name"]
    results = {}

    for fig_key, (ctype, tsuffix) in CHART_FIGURES.items():
        out_path = img_dir / f"{fig_key.replace('-', '_')}.png"
        if out_path.exists():
            results[fig_key] = str(out_path)
            continue

        title = f"{name}{tsuffix or ''}"
        data = gen_chart_data(project, ctype, tsuffix)

        try:
            if ctype == "swot":
                s = [l.strip() for l in data["s"].split("\n") if l.strip()]
                w = [l.strip() for l in data["w"].split("\n") if l.strip()]
                o = [l.strip() for l in data["o"].split("\n") if l.strip()]
                t = [l.strip() for l in data["t"].split("\n") if l.strip()]
                engine.generate_swot_diagram(s, w, o, t, str(out_path))
            elif ctype == "architecture":
                engine.generate_architecture_diagram(title, str(out_path))
            elif ctype == "gantt":
                engine.generate_gantt_diagram(data.get("phases", []), str(out_path))
            elif ctype == "bmc":
                engine.generate_bmc_diagram(data, str(out_path))
            elif ctype == "value_chain":
                engine.generate_value_chain_diagram(
                    data.get("nodes", []), data.get("edges", []), str(out_path))
            elif ctype == "competition":
                engine.generate_competition_map(data.get("players", []), str(out_path))

            results[fig_key] = str(out_path)
        except Exception as e:
            print(f"    {fig_key} FAIL: {e}")

    print(f"  Charts: {len(results)}/{len(CHART_FIGURES)} generated")
    return results


# ---- DOCX图片插入 ----
def insert_images(docx_path, img_dir, screenshots, charts):
    from docx import Document
    from docx.shared import Inches
    from docx.oxml.ns import qn

    doc = Document(str(docx_path))

    # 构建映射: fig_key -> image_path
    image_map = {}
    for fig_key, path in charts.items():
        image_map[fig_key] = path
    if screenshots:
        for fig_key, idx in SCREENSHOT_FIGURES:
            if idx < len(screenshots):
                image_map[fig_key] = screenshots[idx]["path"]

    inserted = 0
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        m = re.match(r"^(图\d+-\d+)\s", p.text.strip())
        if not m:
            continue
        fig_key = m.group(1)
        img_path = image_map.get(fig_key)
        if not img_path or not os.path.exists(img_path):
            continue

        # 找到图注前面的空段落（图片位置）
        target_para = None
        if i > 0 and not paras[i - 1].text.strip():
            target_para = paras[i - 1]

        if target_para is None:
            continue

        # 清除空段落中的残留内容
        for child in list(target_para._element):
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "r":
                target_para._element.remove(child)

        # 根据图片宽高比决定插入尺寸
        try:
            w, h = PILImage.open(img_path).size
            if h > w * 1.5:  # 竖屏图
                width = Inches(2.8)
            else:
                width = Inches(5.2)
        except:
            width = Inches(5.2)

        run = target_para.add_run()
        try:
            run.add_picture(img_path, width=width)
            inserted += 1
        except Exception as e:
            print(f"    Insert {fig_key} fail: {e}")

    # 居中对齐所有含图片的段落
    for p in doc.paragraphs:
        if p.text.strip() == "":
            has_img = False
            for r in p.runs:
                drawings = r._element.findall(qn("w:drawing"))
                if drawings:
                    has_img = True
                    break
            if has_img:
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(docx_path))
    print(f"  DOCX: {inserted} images inserted -> {docx_path.name}")
    return inserted


# ---- 主流程 ----
def main(max_projects=0, start_idx=0):
    print("=" * 60)
    print(" 湖北省创业扶持项目 - 批量图片生成与插入")
    print("=" * 60)

    # 加载进度
    progress = {}
    if PROGRESS_FILE.exists():
        try:
            progress = json.loads(PROGRESS_FILE.read_text(encoding="utf-8"))
        except:
            pass

    # 加载项目列表
    projects = load_projects()
    end_idx = len(projects)
    if max_projects > 0:
        end_idx = min(start_idx + max_projects, len(projects))
    print(f"共 {len(projects)} 个项目, 处理范围: {start_idx}-{end_idx-1}")

    # PPT引擎
    from ppt_diagram_engine import PPTDiagramEngine
    engine = None
    engine_count = 0
    RESTART_EVERY = 5  # 每5个项目重启PPT

    for idx in range(start_idx, end_idx):
        pid = f"p{idx}"
        if progress.get(pid, {}).get("done"):
            print(f"\n[{idx+1}/{len(projects)}] {projects[idx]['name']} - 已完成，跳过")
            continue

        print(f"\n{'─'*50}")
        print(f"[{idx+1}/{len(projects)}] {projects[idx]['name']}")

        docx_path = find_docx_for_project(projects[idx]["name"])
        if not docx_path:
            print("  DOCX未找到，跳过")
            continue

        img_dir = IMAGE_BASE / pid
        img_dir.mkdir(parents=True, exist_ok=True)

        # 阶段1: 截图
        t0 = time.time()
        try:
            screenshots = gen_screenshots(projects[idx], img_dir)
        except Exception as e:
            print(f"  截图生成失败: {e}")
            screenshots = []
        t1 = time.time()

        # 阶段2: 图表 (定期重启PPT)
        if engine is None or engine_count >= RESTART_EVERY:
            if engine is not None:
                try:
                    engine.quit()
                except:
                    pass
                time.sleep(2)
            try:
                engine = PPTDiagramEngine(visible=False)
                engine_count = 0
                print("  PPT引擎已启动")
            except Exception as e:
                print(f"  PPT引擎启动失败: {e}")
                engine = None

        charts = {}
        if engine:
            try:
                charts = gen_charts(engine, projects[idx], img_dir)
                engine_count += 1
            except Exception as e:
                print(f"  图表生成异常: {e}")
                try:
                    engine.quit()
                except:
                    pass
                engine = None
        t2 = time.time()

        # 阶段3: 插入DOCX
        try:
            inserted = insert_images(docx_path, img_dir, screenshots, charts)
        except Exception as e:
            print(f"  DOCX插入失败: {e}")
            inserted = 0
        t3 = time.time()

        # 保存进度
        progress[pid] = {
            "done": True,
            "screenshots": len(screenshots) if screenshots else 0,
            "charts": len(charts),
            "inserted": inserted,
            "time_screenshots": round(t1 - t0, 1),
            "time_charts": round(t2 - t1, 1),
            "time_insert": round(t3 - t2, 1),
        }
        PROGRESS_FILE.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")

        total_imgs = (len(screenshots) if screenshots else 0) + len(charts)
        print(f"  完成: {total_imgs}张图片生成, {inserted}张已插入 "
              f"(截图{t1-t0:.0f}s + 图表{t2-t1:.0f}s + 插入{t3-t2:.0f}s)")

    # 清理
    if engine:
        try:
            engine.quit()
        except:
            pass

    # 汇总
    done = sum(1 for v in progress.values() if v.get("done"))
    total_inserted = sum(v.get("inserted", 0) for v in progress.values())
    print(f"\n{'='*60}")
    print(f"全部完成: {done}/{len(projects)} 个项目, 共插入 {total_inserted} 张图片")
    print(f"{'='*60}")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--max", type=int, default=0, help="只处理前N个项目(0=全部)")
    parser.add_argument("--start", type=int, default=0, help="从第N个项目开始(0-based)")
    args = parser.parse_args()
    main(max_projects=args.max, start_idx=args.start)
