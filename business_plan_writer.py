#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
商业计划书撰写系统 - 完全基于《商业计划书大纲2026》标准
"""

import json
from pathlib import Path
from typing import Dict, List, Optional, Any

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class BusinessPlanWriter:
    """商业计划书撰写器 - 基于《商业计划书大纲2026》"""

    def __init__(self, outline_path: Optional[str] = None):
        """初始化商业计划书撰写器，加载大纲"""
        self.outline_path = outline_path or self._default_outline_path()
        self.outline = self._load_outline()

    def _default_outline_path(self) -> str:
        """获取默认大纲路径"""
        return str(Path(__file__).parent / "outline_with_prompts.json")

    def _load_outline(self) -> Dict[str, Any]:
        """加载大纲文件"""
        try:
            with open(self.outline_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except FileNotFoundError:
            print(f"警告: 大纲文件未找到 {self.outline_path}")
            return self._get_builtin_outline()

    def _get_builtin_outline(self) -> Dict[str, Any]:
        """获取内置大纲（简化版，主要用于备用）"""
        return {
            "template_name": "商业计划书大纲2026",
            "total_sections": 76,
            "sections": []
        }

    def create_project_info(self, product_name: str, company_name: str,
                            project_desc: str = "", **kwargs) -> Dict[str, str]:
        """创建项目信息"""
        return {
            "product_name": product_name,
            "company_name": company_name,
            "project_description": project_desc,
            **kwargs
        }

    def generate_content_plan(self, project_info: Dict[str, str]) -> Dict[str, Any]:
        """生成内容计划（基于大纲）"""
        content_plan = {
            "project_info": project_info,
            "template_name": self.outline.get("template_name", "商业计划书大纲2026"),
            "sections": []
        }

        for sec in self.outline.get("sections", []):
            prompt = self._inject_project_info(sec.get("prompt", ""), project_info)
            
            content_plan["sections"].append({
                "id": sec.get("id"),
                "title": sec.get("title"),
                "level": sec.get("level", 2),
                "type": sec.get("type", "generated"),
                "prompt": prompt,
                "fixed_text": sec.get("fixed_text", ""),
                "needs_chart": sec.get("needs_chart", False),
                "chart_type": sec.get("chart_type"),
                "table_skeleton": sec.get("table_skeleton"),
                "generated_content": ""
            })

        return content_plan

    def _inject_project_info(self, prompt: str, project_info: Dict[str, str]) -> str:
        """将项目信息注入提示词"""
        full_prompt = f"项目名称: {project_info.get('product_name', '未命名项目')}\n"
        full_prompt += f"公司名称: {project_info.get('company_name', '未注册公司')}\n"
        if project_info.get('project_description'):
            full_prompt += f"项目描述: {project_info.get('project_description')}\n"
        full_prompt += "\n" + prompt
        return full_prompt

    def save_content_plan(self, content_plan: Dict[str, Any], output_path: str):
        """保存内容计划到JSON文件"""
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(content_plan, f, ensure_ascii=False, indent=2)
        print(f"内容计划已保存到: {output_path}")

    def load_content_plan(self, input_path: str) -> Dict[str, Any]:
        """从JSON文件加载内容计划"""
        with open(input_path, 'r', encoding='utf-8') as f:
            return json.load(f)

    def get_outline_info(self) -> Dict[str, Any]:
        """获取大纲信息"""
        return {
            "template_name": self.outline.get("template_name", "商业计划书大纲2026"),
            "total_sections": self.outline.get("total_sections", 0),
            "sections_count": len(self.outline.get("sections", [])),
            "sections": self.outline.get("sections", [])
        }


class DocxBuilder:
    """使用python-docx构建Word文档 - 完全基于《商业计划书大纲2026》标准"""

    def __init__(self):
        if not HAS_DOCX:
            raise ImportError("需要安装python-docx库: pip install python-docx")

    def build_document(self, content_plan: Dict[str, Any], output_path: str):
        """构建Word文档"""
        doc = Document()

        self._set_default_font(doc)

        project_info = content_plan.get("project_info", {})
        title_text = f"{project_info.get('product_name', '商业计划书')}"
        title = doc.add_heading(title_text, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if project_info.get('company_name'):
            company_para = doc.add_paragraph()
            company_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            company_run = company_para.add_run(project_info['company_name'])
            company_run.font.size = Pt(14)

        doc.add_paragraph()

        sections = content_plan.get("sections", [])
        for sec in sections:
            self._add_section(doc, sec)

        doc.save(output_path)
        print(f"文档生成完成: {output_path}")

    def _set_default_font(self, doc):
        """设置默认字体 - 符合中文商业计划书标准"""
        doc.styles['Normal'].font.name = '宋体'
        doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        doc.styles['Normal'].font.size = Pt(12)

        for i in range(1, 4):
            style = doc.styles[f'Heading {i}']
            style.font.name = '黑体'
            style._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    def _add_section(self, doc, sec):
        """添加一个章节"""
        title = sec.get("title", "")
        level = sec.get("level", 2)

        if title:
            doc.add_heading(title, level=level)

        sec_type = sec.get("type", "generated")
        if sec_type == "table" and sec.get("table_skeleton"):
            table_data = sec["table_skeleton"].get("data", [])
            if table_data:
                rows = len(table_data)
                cols = max(len(r) for r in table_data) if table_data else 1
                table = doc.add_table(rows=rows, cols=cols)
                table.style = 'Table Grid'

                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        cell = table.rows[r_idx].cells[c_idx]
                        cell.text = str(val)

        text_content = ""
        if sec.get("generated_content"):
            text_content = sec["generated_content"]
        elif sec.get("content"):
            text_content = sec["content"]
        elif sec.get("prompt"):
            text_content = f"[待生成内容: {sec.get('title', '')}]\n\n提示词:\n{sec.get('prompt', '')}"
        elif sec.get("fixed_text"):
            text_content = sec["fixed_text"]

        if text_content and sec_type != "table":
            for para_text in str(text_content).split("\n"):
                para_text = para_text.strip()
                if para_text:
                    para = doc.add_paragraph(para_text)
                    para.paragraph_format.line_spacing = 1.5

        if sec.get("needs_chart") and sec.get("chart_path"):
            try:
                doc.add_picture(sec["chart_path"], width=Inches(6))
            except Exception as e:
                print(f"添加图片失败: {e}")

        doc.add_paragraph()


def main():
    """主函数演示"""
    print("=" * 60)
    print("商业计划书撰写系统 - 《商业计划书大纲2026》标准")
    print("=" * 60)

    writer = BusinessPlanWriter()
    outline_info = writer.get_outline_info()

    print(f"\n大纲模板: {outline_info.get('template_name')}")
    print(f"总章节数: {outline_info.get('total_sections')}")
    print(f"实际章节: {len(outline_info.get('sections', []))}")

    print("\n示例：创建内容计划")
    project_info = writer.create_project_info(
        product_name="示例项目",
        company_name="示例公司",
        project_description="这是一个示例项目"
    )

    content_plan = writer.generate_content_plan(project_info)
    output_dir = Path(__file__).parent / "output"
    output_dir.mkdir(exist_ok=True)
    plan_path = output_dir / "content_plan.json"
    writer.save_content_plan(content_plan, str(plan_path))

    print("\n内容计划已生成！")


if __name__ == "__main__":
    main()
